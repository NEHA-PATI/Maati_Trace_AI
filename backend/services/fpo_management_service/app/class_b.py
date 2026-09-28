from __future__ import annotations

import csv
import hashlib
import io
import json
import re
from datetime import datetime, timedelta, timezone
from pathlib import PurePath
from typing import Any
from uuid import UUID

from sqlalchemy import text
from sqlalchemy.engine import Connection

from shared.db.postgres import engine
from services.fpo_management_service.app.repository import get_fpo_access_context, require_fpo_feature
from services.fpo_management_service.app.bulk_states import staged_status_for_farmer_match
from services.fpo_management_service.app.storage import download_bytes, store_bytes


DEFAULT_IMPORT_CONFIG = {
    "max_rows_per_job": 5000,
    "max_jobs_per_day": 10,
    "allowed_formats": ["CSV", "XLSX"],
    "require_dry_run": True,
    "artifact_retention_days": 30,
}
FORMULA_PREFIXES = ("=", "+", "-", "@")
FORBIDDEN_IMPORT_FIELDS = {"farmer_id", "user_id", "password", "password_hash", "otp", "role"}


def _owner(conn: Connection, user_id: UUID | str, feature: str) -> tuple[dict[str, Any], dict[str, Any]]:
    context = get_fpo_access_context(conn, user_id=user_id)
    entitlements = require_fpo_feature(conn, user_id=user_id, feature_key=feature)
    effective = entitlements["features"][feature]
    return context, effective


def _safe_filename(filename: str) -> str:
    value = PurePath(filename).name.strip()
    value = re.sub(r"[^A-Za-z0-9._-]", "_", value)
    if not value or value.startswith("."):
        raise ValueError("A safe filename is required")
    return value[:180]


def _config(feature: dict[str, Any]) -> dict[str, Any]:
    value = dict(DEFAULT_IMPORT_CONFIG)
    value.update(feature.get("configuration") or {})
    return value


def import_template(conn: Connection, *, user_id: UUID | str, import_type: str) -> dict[str, Any]:
    _owner(conn, user_id, "BULK_FARM_REGISTRATION")
    kind = import_type.upper()
    if kind not in {"FARMERS", "FARMS", "FARMERS_AND_FARMS"}:
        raise ValueError("Unsupported import type")
    columns = {
        "FARMERS": ["external_reference", "farmer_name", "phone", "email", "state_name", "district_name", "block_name", "village_name"],
        "FARMS": ["farmer_external_reference", "farm_external_reference", "farm_name", "state_name", "district_name", "block_name", "village_name", "crop_name", "area_acres", "land_record_reference"],
        "FARMERS_AND_FARMS": ["external_reference", "farmer_name", "phone", "email", "state_name", "district_name", "block_name", "village_name", "farm_external_reference", "farm_name", "crop_name", "area_acres", "land_record_reference"],
    }[kind]
    return {
        "import_type": kind,
        "template_version": 2,
        "format": "CSV",
        "columns": columns,
        "field_help": {
            "external_reference": "Your stable farmer/member reference. Not a MaatiTrace farmer ID.",
            "farmer_external_reference": "Your existing farmer/member reference used to match a farmer.",
            "farm_external_reference": "Your stable reference for this farm/land parcel.",
            "land_record_reference": "Optional official land-record or survey reference; not a MaatiTrace ID.",
            "crop_name": "Human-readable crop name. The canonical crop code is resolved automatically.",
            "crop_code": "Optional canonical code for integrations; crop_name is preferred for manual uploads.",
        },
        "rules": {
            "formula_cells": "rejected",
            "dry_run_required": True,
            "farmer_id": "never_entered_in_csv",
            "crop_code": "optional_when_crop_name_is_present",
        },
    }


def _resolve_crop(conn: Connection, payload: dict[str, Any]) -> tuple[str | None, str | None, list[dict[str, str]]]:
    crop_code = (payload.get("crop_code") or "").strip().lower() or None
    crop_name = (payload.get("crop_name") or "").strip() or None
    if not crop_code and not crop_name:
        return None, None, [{"field": "crop_name", "code": "REQUIRED", "message": "Enter a crop name; crop code is optional."}]
    row = conn.execute(text("""
        SELECT c.crop_code, COALESCE(t.display_name, c.crop_code) AS crop_name
        FROM crop_observation.crops c
        LEFT JOIN crop_observation.crop_translations t
          ON t.crop_id = c.crop_id AND t.locale IN ('en', 'en-IN')
        WHERE c.is_active = TRUE
          AND (:crop_code IS NULL OR lower(c.crop_code) = :crop_code)
          AND (:crop_name IS NULL OR lower(COALESCE(t.display_name, c.crop_code)) = lower(:crop_name))
        ORDER BY CASE WHEN lower(c.crop_code) = :crop_code THEN 0 ELSE 1 END
        LIMIT 1
    """), {"crop_code": crop_code, "crop_name": crop_name}).mappings().first()
    if not row:
        return None, crop_name, [{"field": "crop_name", "code": "UNKNOWN_CROP", "message": "Crop name does not match an active crop catalogue entry."}]
    return row["crop_code"], row["crop_name"], []


def _normalize_import_row(conn: Connection, payload: dict[str, Any], import_type: str) -> tuple[dict[str, Any], list[dict[str, str]]]:
    normalized = dict(payload)
    errors = []
    for field in FORBIDDEN_IMPORT_FIELDS:
        if normalized.get(field):
            errors.append({"field": field, "code": "FORBIDDEN_FIELD", "message": f"{field} is generated by MaatiTrace and must not be uploaded."})
    if normalized.get("external_reference") and not normalized.get("farmer_external_reference"):
        normalized["farmer_external_reference"] = normalized["external_reference"]
    if import_type in {"FARMS", "FARMERS_AND_FARMS"}:
        code, name, crop_errors = _resolve_crop(conn, normalized)
        normalized["crop_code"] = code
        normalized["crop_name"] = name or normalized.get("crop_name")
        errors.extend(crop_errors)
    return normalized, errors


def create_import_intent(conn: Connection, *, user_id: UUID | str, import_type: str, source_format: str, filename: str, checksum: str) -> dict[str, Any]:
    context, feature = _owner(conn, user_id, "BULK_FARM_REGISTRATION")
    kind, fmt = import_type.upper(), source_format.upper()
    config = _config(feature["features"]["BULK_FARM_REGISTRATION"] if "features" in feature else feature)
    if kind not in {"FARMERS", "FARMS", "FARMERS_AND_FARMS"} or fmt not in config["allowed_formats"]:
        raise ValueError("Import type or format is not enabled")
    count = conn.execute(text("SELECT count(*) FROM public.fpo_bulk_import_jobs WHERE fpo_id=:fpo AND created_at >= date_trunc('day', now()) AND status <> 'CANCELLED'"), {"fpo": str(context["fpo_id"])}).scalar_one()
    if int(count) >= int(config["max_jobs_per_day"]):
        raise ValueError("Daily import limit has been reached")
    safe_name = _safe_filename(filename)
    row = conn.execute(text("""
        INSERT INTO public.fpo_bulk_import_jobs (fpo_id, import_type, source_format, original_filename, object_key, file_checksum, template_version, configuration_snapshot, requested_by)
        VALUES (:fpo,:kind,:fmt,:filename,:object_key,:checksum,2,CAST(:config AS jsonb),:user)
        RETURNING import_job_id, fpo_id, import_type, source_format, original_filename, file_checksum, template_version, mode, status, configuration_snapshot, expires_at
    """), {"fpo": str(context["fpo_id"]), "kind": kind, "fmt": fmt, "filename": safe_name, "object_key": f"fpo-imports/{context['fpo_id']}/{UUID(int=0)}/{safe_name}", "checksum": checksum.lower(), "config": json.dumps(config), "user": str(user_id)}).mappings().one()
    object_key = f"fpo-imports/{context['fpo_id']}/{row['import_job_id']}/{safe_name}"
    conn.execute(text("UPDATE public.fpo_bulk_import_jobs SET object_key=:key WHERE import_job_id=:id"), {"key": object_key, "id": str(row["import_job_id"])})
    result = dict(row); result["object_key"] = object_key; result["upload_url"] = f"/v1/fpo/me/imports/{row['import_job_id']}/content"; result["method"] = "PUT"; return result


def upload_import_content(conn: Connection, *, user_id: UUID | str, job_id: UUID | str, data: bytes) -> dict[str, Any]:
    row = conn.execute(text("SELECT * FROM public.fpo_bulk_import_jobs j JOIN public.fpo_organizations o ON o.fpo_id=j.fpo_id WHERE j.import_job_id=:id AND o.auth_user_id=:user FOR UPDATE"), {"id": str(job_id), "user": str(user_id)}).mappings().first()
    if not row: raise ValueError("Import job was not found")
    checksum = hashlib.sha256(data).hexdigest()
    if checksum != str(row["file_checksum"]).lower(): raise ValueError("Uploaded file checksum does not match the upload intent")
    store_bytes(object_key=row["object_key"], data=data)
    conn.execute(text("UPDATE public.fpo_bulk_import_jobs SET status='QUEUED', updated_at=now() WHERE import_job_id=:id"), {"id": str(job_id)})
    return {"import_job_id": job_id, "status": "QUEUED", "checksum": checksum, "size_bytes": len(data)}


def _parse_csv(data: bytes) -> list[dict[str, Any]]:
    try: text_data = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc: raise ValueError("CSV must be UTF-8 encoded") from exc
    reader = csv.DictReader(io.StringIO(text_data))
    if not reader.fieldnames: raise ValueError("The file has no header row")
    headers = [str(value or "").strip().lower() for value in reader.fieldnames]
    forbidden_headers = sorted(set(headers) & FORBIDDEN_IMPORT_FIELDS)
    if forbidden_headers:
        raise ValueError(f"Internal fields are not allowed in the CSV header: {', '.join(forbidden_headers)}")
    rows=[]
    for number, raw in enumerate(reader, start=2):
        payload={str(k).strip().lower(): (v or "").strip() for k,v in raw.items() if k}
        errors=[]
        for key,value in payload.items():
            if value.startswith(FORMULA_PREFIXES): errors.append({"field":key,"code":"FORMULA_CELL","message":"Formula-like cells are not accepted"})
        rows.append((number,payload,errors))
    return rows


def _parse_xlsx(data: bytes) -> list[tuple[int, dict[str, Any], list[dict[str, str]]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError("XLSX support is not installed on this worker") from exc
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
    if len(workbook.sheetnames) != 1:
        raise ValueError("The workbook must contain exactly one worksheet")
    sheet = workbook.active
    values = list(sheet.iter_rows(values_only=True))
    if not values:
        raise ValueError("The workbook has no header row")
    headers = [str(value or "").strip().lower() for value in values[0]]
    if not any(headers):
        raise ValueError("The workbook has no header row")
    forbidden_headers = sorted(set(headers) & FORBIDDEN_IMPORT_FIELDS)
    if forbidden_headers:
        raise ValueError(f"Internal fields are not allowed in the XLSX header: {', '.join(forbidden_headers)}")
    result = []
    for number, cells in enumerate(values[1:], start=2):
        payload = {header: str(cells[index] or "").strip() for index, header in enumerate(headers) if header}
        errors = [{"field": key, "code": "FORMULA_CELL", "message": "Formula-like cells are not accepted"} for key, value in payload.items() if value.startswith(FORMULA_PREFIXES)]
        result.append((number, payload, errors))
    return result


def validate_import(conn: Connection, *, user_id: UUID | str, job_id: UUID | str) -> dict[str, Any]:
    row=conn.execute(text("SELECT j.* FROM public.fpo_bulk_import_jobs j JOIN public.fpo_organizations o ON o.fpo_id=j.fpo_id WHERE j.import_job_id=:id AND o.auth_user_id=:user FOR UPDATE"), {"id":str(job_id),"user":str(user_id)}).mappings().first()
    if not row: raise ValueError("Import job was not found")
    data=download_bytes(object_key=row["object_key"])
    if hashlib.sha256(data).hexdigest()!=str(row["file_checksum"]).lower(): raise ValueError("Stored file checksum mismatch")
    parsed = _parse_csv(data) if row["source_format"] == "CSV" else _parse_xlsx(data)
    required={"FARMERS":{"farmer_name"},"FARMS":{"farmer_external_reference","farm_name"},"FARMERS_AND_FARMS":{"farmer_name","farm_name"}}[row["import_type"]]
    conn.execute(text("DELETE FROM public.fpo_bulk_import_rows WHERE import_job_id=:id"), {"id":str(job_id)})
    valid=invalid=0
    for number,payload,errors in parsed:
        payload, normalization_errors = _normalize_import_row(conn, payload, row["import_type"])
        errors.extend(normalization_errors)
        errors += [{"field": field, "code":"REQUIRED", "message": "Required value is missing"} for field in required if not payload.get(field)]
        status="INVALID" if errors else "VALID"; valid += status=="VALID"; invalid += status=="INVALID"
        idem=hashlib.sha256(f"{job_id}:{number}:{json.dumps(payload,sort_keys=True)}".encode()).hexdigest()
        conn.execute(text("INSERT INTO public.fpo_bulk_import_rows (import_job_id,row_number,external_reference,raw_payload,normalized_payload,row_status,error_items,idempotency_key) VALUES (:job,:num,:external,CAST(:raw AS jsonb),CAST(:norm AS jsonb),:status,CAST(:errors AS jsonb),:idem)"), {"job":str(job_id),"num":number,"external":payload.get("farmer_external_reference") or payload.get("external_reference") or payload.get("farm_external_reference"),"raw":json.dumps(payload),"norm":json.dumps(payload) if not errors else None,"status":status,"errors":json.dumps(errors),"idem":idem})
    conn.execute(text("UPDATE public.fpo_bulk_import_jobs SET status='DRY_RUN_READY', total_rows=:total, valid_rows=:valid, invalid_rows=:invalid, mode='DRY_RUN', updated_at=now() WHERE import_job_id=:id"), {"id":str(job_id),"total":len(parsed),"valid":valid,"invalid":invalid})
    return {"import_job_id":job_id,"status":"DRY_RUN_READY","total_rows":len(parsed),"valid_rows":valid,"invalid_rows":invalid}


def list_imports(conn: Connection, *, user_id: UUID | str) -> list[dict[str, Any]]:
    context,_=_owner(conn,user_id,"BULK_FARM_REGISTRATION")
    return [dict(x) for x in conn.execute(text("SELECT import_job_id,import_type,source_format,original_filename,file_checksum,template_version,mode,status,total_rows,valid_rows,invalid_rows,staged_rows,created_rows,updated_rows,skipped_rows,failed_rows,started_at,completed_at,expires_at,error_code,error_message,created_at,updated_at FROM public.fpo_bulk_import_jobs WHERE fpo_id=:fpo ORDER BY created_at DESC LIMIT 100"),{"fpo":str(context["fpo_id"])}).mappings()]


def get_import(conn: Connection, *, user_id: UUID | str, job_id: UUID | str) -> dict[str, Any]:
    context,_=_owner(conn,user_id,"BULK_FARM_REGISTRATION"); row=conn.execute(text("SELECT * FROM public.fpo_bulk_import_jobs WHERE import_job_id=:id AND fpo_id=:fpo"),{"id":str(job_id),"fpo":str(context["fpo_id"])}).mappings().first()
    if not row: raise ValueError("Import job was not found")
    return dict(row)


def import_rows(conn: Connection, *, user_id: UUID | str, job_id: UUID | str) -> list[dict[str, Any]]:
    job=get_import(conn,user_id=user_id,job_id=job_id)
    return [dict(x) for x in conn.execute(text("SELECT import_row_id,row_number,external_reference,normalized_payload,row_status,error_items,farmer_id,farm_id,idempotency_key,processed_at FROM public.fpo_bulk_import_rows WHERE import_job_id=:id ORDER BY row_number"),{"id":str(job["import_job_id"])}).mappings()]


def staged_import_records(conn: Connection, *, user_id: UUID | str, job_id: UUID | str) -> list[dict[str, Any]]:
    job = get_import(conn, user_id=user_id, job_id=job_id)
    return [dict(row) for row in conn.execute(text("""
        SELECT staged_record_id, import_row_id, record_type, farmer_id, farm_id,
               payload, status, status_reason, created_at, updated_at
        FROM public.fpo_bulk_staged_records
        WHERE import_job_id = :job
        ORDER BY created_at, staged_record_id
    """), {"job": str(job["import_job_id"])}).mappings()]


def reconcile_farmer_account(conn: Connection, *, payload: dict[str, Any]) -> int:
    """Attach a newly-created farmer account without granting FPO access."""
    user_id = payload.get("user_id")
    email = (payload.get("email") or "").strip().lower() or None
    phone = (payload.get("phone_number") or "").strip() or None
    if not user_id or not (email or phone):
        return 0
    farmer_id = conn.execute(text("""
        SELECT fp.farmer_id FROM public.farmer_profiles fp
        WHERE fp.user_id = :user_id AND fp.is_active = TRUE LIMIT 1
    """), {"user_id": str(user_id)}).scalar_one_or_none()
    if not farmer_id:
        return 0
    result = conn.execute(text("""
        UPDATE public.fpo_bulk_staged_records s
        SET farmer_id = :farmer_id,
            status = CASE WHEN s.record_type = 'FARM' THEN 'AWAITING_FARM_CONFIRMATION' ELSE 'AWAITING_CONSENT' END,
            status_reason = CASE WHEN s.record_type = 'FARM'
                THEN 'Farmer account matched. Farmer confirmation and consent are still required.'
                ELSE 'Farmer account matched. Consent is still required.' END,
            updated_at = now()
        WHERE s.farmer_id IS NULL
          AND (lower(COALESCE(s.payload->>'email', '')) = :email
               OR COALESCE(s.payload->>'phone', '') = :phone)
    """), {"farmer_id": str(farmer_id), "email": email or "", "phone": phone or ""})
    return int(result.rowcount or 0)


def farmer_staged_records(conn: Connection, *, user_id: UUID | str) -> list[dict[str, Any]]:
    rows = conn.execute(text("""
        SELECT s.staged_record_id, s.import_job_id, s.record_type, s.payload,
               s.status, s.status_reason, s.created_at, s.updated_at
        FROM public.fpo_bulk_staged_records s
        JOIN public.farmer_profiles fp ON fp.farmer_id = s.farmer_id
        WHERE fp.user_id = :user_id AND fp.is_active = TRUE
        ORDER BY s.created_at DESC
    """), {"user_id": str(user_id)}).mappings()
    return [dict(row) for row in rows]


def farmer_staged_record(conn: Connection, *, user_id: UUID | str, staged_record_id: UUID | str) -> dict[str, Any]:
    row = conn.execute(text("""
        SELECT s.staged_record_id, s.import_job_id, s.record_type, s.farmer_id,
               s.farm_id, s.payload, s.status, s.status_reason, s.created_at, s.updated_at
        FROM public.fpo_bulk_staged_records s
        JOIN public.farmer_profiles fp ON fp.farmer_id = s.farmer_id
        WHERE s.staged_record_id = :staged AND fp.user_id = :user_id AND fp.is_active = TRUE
        LIMIT 1
    """), {"staged": str(staged_record_id), "user_id": str(user_id)}).mappings().first()
    if not row:
        raise ValueError("Imported onboarding record was not found")
    return dict(row)


def link_staged_farm(conn: Connection, *, user_id: UUID | str, staged_record_id: UUID | str, farm_id: UUID | str) -> dict[str, Any]:
    row = farmer_staged_record(conn, user_id=user_id, staged_record_id=staged_record_id)
    if row["record_type"] != "FARM":
        raise ValueError("This onboarding record does not contain a farm")
    farm = conn.execute(text("""
        SELECT farm_id FROM public.farms
        WHERE farm_id = :farm AND farmer_id = :farmer AND is_active = TRUE
        LIMIT 1
    """), {"farm": str(farm_id), "farmer": str(row["farmer_id"])}).scalar_one_or_none()
    if not farm:
        raise ValueError("The farm does not belong to the authenticated farmer")
    updated = conn.execute(text("""
        UPDATE public.fpo_bulk_staged_records
        SET farm_id = :farm, status = 'AWAITING_CONSENT',
            status_reason = 'Farm confirmed. Review the FPO request and consent before sharing data.', updated_at = now()
        WHERE staged_record_id = :staged
        RETURNING staged_record_id, farm_id, status, status_reason, updated_at
    """), {"farm": str(farm_id), "staged": str(staged_record_id)}).mappings().one()
    conn.execute(text("""
        UPDATE public.fpo_bulk_import_rows SET farm_id = :farm
        WHERE import_row_id = (SELECT import_row_id FROM public.fpo_bulk_staged_records WHERE staged_record_id = :staged)
    """), {"farm": str(farm_id), "staged": str(staged_record_id)})
    return dict(updated)


def commit_import(conn: Connection, *, user_id: UUID | str, job_id: UUID | str) -> dict[str, Any]:
    job=get_import(conn,user_id=user_id,job_id=job_id)
    if job["status"]!="DRY_RUN_READY": raise ValueError("A successful dry run is required before commit")
    if int(job["invalid_rows"]): raise ValueError("Correct invalid rows before commit")
    conn.execute(text("UPDATE public.fpo_bulk_import_jobs SET status='COMMITTING', mode='COMMIT', started_at=now(), updated_at=now() WHERE import_job_id=:id AND version=:version"),{"id":str(job_id),"version":job["version"]})
    return {"import_job_id":job_id,"status":"COMMITTING","message":"Queued for staged farmer onboarding. Account creation, farm confirmation, and consent remain explicit user actions."}


def _find_existing_farmer(conn: Connection, payload: dict[str, Any]) -> dict[str, Any] | None:
    email = (payload.get("email") or "").strip().lower() or None
    phone = (payload.get("phone") or "").strip() or None
    if email:
        row = conn.execute(text("""
            SELECT fp.farmer_id
            FROM public.farmer_profiles fp
            JOIN public.users u ON u.user_id = fp.user_id
            WHERE u.role = 'farmer' AND lower(u.email) = :email AND fp.is_active = TRUE
            LIMIT 1
        """), {"email": email}).mappings().first()
        if row:
            return dict(row)
    if phone:
        row = conn.execute(text("""
            SELECT fp.farmer_id
            FROM public.farmer_profiles fp
            JOIN public.users u ON u.user_id = fp.user_id
            WHERE u.role = 'farmer' AND u.phone_number = :phone AND fp.is_active = TRUE
            LIMIT 1
        """), {"phone": phone}).mappings().first()
        if row:
            return dict(row)
    return None


def process_import_jobs(batch_size: int = 10) -> int:
    processed = 0
    with engine.begin() as conn:
        jobs = list(conn.execute(text("""
            SELECT import_job_id, fpo_id, import_type
            FROM public.fpo_bulk_import_jobs
            WHERE status = 'COMMITTING'
            ORDER BY created_at
            FOR UPDATE SKIP LOCKED
            LIMIT :limit
        """), {"limit": batch_size}).mappings())
        for job in jobs:
            rows = list(conn.execute(text("""
                SELECT import_row_id, normalized_payload
                FROM public.fpo_bulk_import_rows
                WHERE import_job_id = :job AND row_status = 'VALID'
                ORDER BY row_number
            """), {"job": str(job["import_job_id"])}).mappings())
            staged = 0
            for row in rows:
                payload = row["normalized_payload"] or {}
                farmer = _find_existing_farmer(conn, payload)
                record_type = "FARM" if job["import_type"] in {"FARMS", "FARMERS_AND_FARMS"} else "FARMER"
                if farmer:
                    status = "AWAITING_FARM_CONFIRMATION" if record_type == "FARM" else "AWAITING_CONSENT"
                    reason = "Existing farmer matched; farmer confirmation and consent are still required."
                    farmer_id = str(farmer["farmer_id"])
                else:
                    status = "AWAITING_FARMER_ACCOUNT"
                    reason = "No existing farmer account matched. Farmer must complete OTP signup before this record can proceed."
                    farmer_id = None
                conn.execute(text("""
                    INSERT INTO public.fpo_bulk_staged_records
                        (fpo_id, import_job_id, import_row_id, record_type, farmer_id, payload, status, status_reason)
                    VALUES (:fpo, :job, :row, :type, :farmer, CAST(:payload AS jsonb), :status, :reason)
                    ON CONFLICT (import_row_id) DO UPDATE SET
                        farmer_id = EXCLUDED.farmer_id,
                        payload = EXCLUDED.payload,
                        status = EXCLUDED.status,
                        status_reason = EXCLUDED.status_reason,
                        updated_at = now()
                """), {"fpo": str(job["fpo_id"]), "job": str(job["import_job_id"]), "row": str(row["import_row_id"]), "type": record_type, "farmer": farmer_id, "payload": json.dumps(payload), "status": status, "reason": reason})
                conn.execute(text("""
                    UPDATE public.fpo_bulk_import_rows
                    SET row_status = 'STAGED', farmer_id = :farmer, processed_at = now()
                    WHERE import_row_id = :row
                """), {"farmer": farmer_id, "row": str(row["import_row_id"])})
                staged += 1
            conn.execute(text("""
                UPDATE public.fpo_bulk_import_jobs
                SET status = 'COMPLETED', staged_rows = :staged, completed_at = now(), updated_at = now()
                WHERE import_job_id = :job
            """), {"staged": staged, "job": str(job["import_job_id"])})
            processed += 1
    return processed


def cancel_import(conn: Connection, *, user_id: UUID | str, job_id: UUID | str) -> dict[str, Any]:
    job=get_import(conn,user_id=user_id,job_id=job_id)
    if job["status"] in {"COMPLETED","CANCELLED"}: raise ValueError("Import is already terminal")
    conn.execute(text("UPDATE public.fpo_bulk_import_jobs SET status='CANCELLED', updated_at=now() WHERE import_job_id=:id"),{"id":str(job_id)})
    return {"import_job_id":job_id,"status":"CANCELLED"}


def quality_summary(conn: Connection, *, user_id: UUID | str) -> dict[str, Any]:
    context,_=_owner(conn,user_id,"DATA_QUALITY_WORKBENCH")
    rows=conn.execute(text("SELECT severity,status,count(*) AS count FROM public.fpo_data_quality_issues WHERE fpo_id=:fpo GROUP BY severity,status ORDER BY severity,status"),{"fpo":str(context["fpo_id"])}).mappings().all()
    return {"fpo_id":context["fpo_id"],"items":[dict(x) for x in rows],"generated_at":datetime.now(timezone.utc)}


def quality_issues(conn: Connection, *, user_id: UUID | str) -> list[dict[str, Any]]:
    context,_=_owner(conn,user_id,"DATA_QUALITY_WORKBENCH")
    return [dict(x) for x in conn.execute(text("SELECT * FROM public.fpo_data_quality_issues WHERE fpo_id=:fpo ORDER BY CASE severity WHEN 'CRITICAL' THEN 1 WHEN 'WARNING' THEN 2 ELSE 3 END,last_detected_at DESC LIMIT 500"),{"fpo":str(context["fpo_id"])}).mappings()]
